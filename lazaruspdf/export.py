"""Export the edited document: PDF, Word (.docx), EPUB, Markdown, HTML, plain text.

Filled gaps export as plain text. Open gaps keep their highlighted [[FILL IN - p.N]] marker.
Images go into PDF/DOCX/EPUB/HTML; TXT/MD keep the [[FIGURE - p.N]] marker."""
import html
import io
import re
import uuid
from pathlib import Path

from .extract import GAP, Figure, Filled
from .output import (FIG_DPI, fig_marker, gap_marker, para_html, render_html, render_pdf,
                     render_text, spaced)

FORMATS = [  # key, menu label, extension, file-dialog filter
    ("pdf", "PDF", ".pdf", "PDF (*.pdf)"),
    ("docx", "Word (.docx)", ".docx", "Word document (*.docx)"),
    ("epub", "EPUB", ".epub", "EPUB e-book (*.epub)"),
    ("md", "Markdown", ".md", "Markdown (*.md)"),
    ("html", "HTML", ".html", "HTML (*.html)"),
    ("txt", "Plain text", ".txt", "Text (*.txt)"),
]

_XML_BAD = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f￾￿]")


def _xml_safe(s):
    return _XML_BAD.sub("", s)


def to_txt(pages, path, title, page_headings=False):
    Path(path).write_text(render_text(pages, page_headings), encoding="utf-8")


def to_html(pages, path, title, page_headings=False):
    Path(path).write_text(render_html(pages, title, clean=True, page_headings=page_headings),
                          encoding="utf-8")


def to_pdf(pages, path, title, page_headings=False):
    import pymupdf
    data = render_pdf(pages, clean=True, page_headings=page_headings)
    doc = pymupdf.open("pdf", data)
    doc.set_metadata({"title": title, "producer": "Lazarus PDF", "creator": "Lazarus PDF"})
    doc.save(str(path), garbage=3, deflate=True)
    doc.close()


# ---------- Markdown ----------

_MD_ESC = re.compile(r"([\\`*_\[\]<>|])")
_MD_LEAD = re.compile(r"^(\s*)([#>+\-=]|\d+[.)])")


def _md_text(s):
    s = _MD_ESC.sub(r"\\\1", s)
    return _MD_LEAD.sub(lambda m: m.group(1) + "\\" + m.group(2), s)


def to_md(pages, path, title, page_headings=False):
    out = []
    for pr in pages:
        if page_headings:
            out += [f"###### Page {pr.number}", ""]
        for it in pr.items:
            if isinstance(it, Figure):
                out.append(f"**{fig_marker(pr.number)}**")
                for lab in it.labels:
                    out.append(f"*Text in figure: {_md_text(lab)}*")
            else:
                parts = []
                for s in spaced(it.segs):
                    if s is GAP:
                        parts.append(f"**{gap_marker(it.page)}**")
                    elif isinstance(s, Filled):
                        parts.append(_md_text(s.text))
                    else:
                        parts.append(_md_text(s))
                out.append("".join(parts))
            out.append("")
    Path(path).write_text("\n".join(out), encoding="utf-8")


# ---------- Word ----------

def to_docx(pages, path, title, page_headings=False):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    doc.core_properties.title = title
    doc.core_properties.author = ""
    doc.styles["Normal"].font.size = Pt(11)

    def marker(par, text, color):
        r = par.add_run(text)
        r.bold = True
        r.font.highlight_color = color
        if color == WD_COLOR_INDEX.YELLOW:
            r.font.color.rgb = RGBColor(0xB7, 0x1C, 0x1C)

    for pr in pages:
        if page_headings:
            h = doc.add_paragraph()
            r = h.add_run(f"Page {pr.number}")
            r.font.size = Pt(8)
            r.font.color.rgb = RGBColor(0x77, 0x77, 0x77)
        for it in pr.items:
            if isinstance(it, Figure):
                marker(doc.add_paragraph(), fig_marker(pr.number), WD_COLOR_INDEX.TURQUOISE)
                from PIL import Image
                px = Image.open(io.BytesIO(it.png)).size[0]
                doc.add_picture(io.BytesIO(it.png), width=Inches(min(6.0, px / FIG_DPI)))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                for lab in it.labels:
                    r = doc.add_paragraph().add_run(_xml_safe(f"Text in figure: {lab}"))
                    r.italic = True
                    r.font.size = Pt(9)
                continue
            p = doc.add_paragraph()
            for s in spaced(it.segs):
                if s is GAP:
                    marker(p, gap_marker(it.page), WD_COLOR_INDEX.YELLOW)
                else:
                    p.add_run(_xml_safe(s.text if isinstance(s, Filled) else s))
    doc.save(str(path))


# ---------- EPUB ----------

_EPUB_CSS = """
body{font-family:serif;line-height:1.5;margin:0 4%}
p{margin:0 0 .8em;text-indent:0}
h2.pg{font-family:sans-serif;font-size:.7em;color:#888;border-top:1px solid #ccc;margin:1.5em 0 .5em}
.gap{background-color:#ffeb3b;color:#b71c1c;font-weight:bold}
.fig{background-color:#4fc3f7;color:#000;font-weight:bold}
figure{margin:1em 0;text-align:center}
figure img{max-width:100%;height:auto}
figcaption{font-size:.8em;color:#666}
"""

PAGES_PER_SECTION = 20


def to_epub(pages, path, title, page_headings=False):
    from ebooklib import epub

    book = epub.EpubBook()
    book.set_identifier(f"urn:uuid:{uuid.uuid4()}")
    book.set_title(title)
    book.set_language("en")
    css = epub.EpubItem(uid="style", file_name="style/main.css", media_type="text/css",
                        content=_EPUB_CSS.encode())
    book.add_item(css)

    sections, n_img = [], 0
    for start in range(0, len(pages), PAGES_PER_SECTION):
        chunk = pages[start:start + PAGES_PER_SECTION]
        body = []
        for pr in chunk:
            if page_headings:
                body.append(f'<h2 class="pg" id="p{pr.number}">Page {pr.number}</h2>')
            for it in pr.items:
                if isinstance(it, Figure):
                    n_img += 1
                    name = f"images/fig{n_img}.png"
                    book.add_item(epub.EpubImage(uid=f"fig{n_img}", file_name=name,
                                                 media_type="image/png", content=it.png))
                    cap = "".join(f"<div>Text in figure: {html.escape(_xml_safe(l))}</div>"
                                  for l in it.labels)
                    body.append(f'<figure><p><span class="fig">{html.escape(fig_marker(pr.number))}</span></p>'
                                f'<img src="{name}" alt="Figure, page {pr.number}"/>'
                                + (f"<figcaption>{cap}</figcaption>" if cap else "") + "</figure>")
                else:
                    body.append(_xml_safe(para_html(it, clean=True)))
        first, last = chunk[0].number, chunk[-1].number
        label = f"Pages {first}–{last}" if first != last else f"Page {first}"
        sec = epub.EpubHtml(title=label, file_name=f"part{len(sections) + 1:03d}.xhtml", lang="en")
        sec.content = "".join(body) or "<p></p>"
        sec.add_item(css)
        book.add_item(sec)
        sections.append(sec)

    book.toc = sections
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    book.spine = ["nav"] + sections
    epub.write_epub(str(path), book)


EXPORTERS = {"pdf": to_pdf, "docx": to_docx, "epub": to_epub, "md": to_md,
             "html": to_html, "txt": to_txt}


def export(fmt, pages, path, title, page_headings=False):
    from . import protect
    protect.check(path)                     # never write over an opened original
    EXPORTERS[fmt](pages, path, title, page_headings)
